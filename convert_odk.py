import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
import re

# --- Lecture du fichier CSV ---
df = pd.read_csv(
    "/root/.claude/uploads/5a50d722-e734-53cc-9ba5-8ea4c49a6d17/816eabe3-adkoul_menage_ciblage.csv",
    dtype=str
)

# --- Traduction des en-têtes ---
col_labels = {
    "SubmissionDate": "Date de soumission",
    "start": "Début",
    "end": "Fin",
    "today": "Date",
    "deviceid": "ID Appareil",
    "username": "Utilisateur",
    "a_village-code_projet": "Code Projet",
    "a_village-region": "Région",
    "a_village-departement": "Département",
    "a_village-commune": "Commune",
    "a_village-note_auteur": "Note Auteur",
    "a_village-village": "Village",
    "a_village-code_village": "Code Village",
    "a_village-seuil_terre_p_zone": "Seuil Terre Pauvre (ha)",
    "a_village-seuil_terre_m_zone": "Seuil Terre Moyen (ha)",
    "a_village-gps_menage-Latitude": "GPS Latitude",
    "a_village-gps_menage-Longitude": "GPS Longitude",
    "a_village-gps_menage-Altitude": "GPS Altitude",
    "a_village-gps_menage-Accuracy": "GPS Précision",
    "b_chef-numero_menage": "N° Ménage",
    "b_chef-cm_nom": "Nom Chef de Ménage",
    "b_chef-cm_sexe": "Sexe",
    "b_chef-cm_age": "Âge",
    "b_chef-cm_tel": "Téléphone",
    "b_chef-cm_cni": "CNI",
    "b_chef-cm_matrimonial": "Statut Matrimonial",
    "c_composition-note_compo": "Note Composition",
    "c_composition-f_0_4": "Femmes 0-4 ans",
    "c_composition-f_5_17": "Femmes 5-17 ans",
    "c_composition-f_18_49": "Femmes 18-49 ans",
    "c_composition-f_50": "Femmes 50+ ans",
    "c_composition-h_0_4": "Hommes 0-4 ans",
    "c_composition-h_5_17": "Hommes 5-17 ans",
    "c_composition-h_18_49": "Hommes 18-49 ans",
    "c_composition-h_50": "Hommes 50+ ans",
    "c_composition-taille_menage": "Taille Ménage",
    "c_composition-note_taille": "Note Taille",
    "c_composition-nb_enfants": "Nb Enfants",
    "c_composition-nb_actifs": "Nb Actifs",
    "c_composition-nb_ages": "Nb Âgés",
    "c_composition-ratio_dependance": "Ratio Dépendance",
    "c_composition-alerte_menage_seul": "Alerte Ménage Seul",
    "d_statut-statut_menage": "Statut Ménage",
    "d_statut-anciennete_depl": "Ancienneté Déplacement",
    "e_protection-nb_plw": "Nb PLW",
    "e_protection-nb_handicap": "Nb Handicapés",
    "e_protection-nb_malade_chronique": "Nb Malades Chroniques",
    "e_protection-nb_age_non_auto": "Nb Âgés Non Autonomes",
    "e_protection-nb_orphelins": "Nb Orphelins",
    "f_categorie-note_cat": "Note Catégorie",
    "f_categorie-cat_communautaire": "Catégorie Communautaire",
    "g_moyens-source_principale": "Source Principale Revenus",
    "g_moyens-source_autre": "Autre Source",
    "g_moyens-nb_actifs_revenu": "Nb Actifs avec Revenus",
    "g_moyens-revenu_mensuel": "Revenu Mensuel (FCFA)",
    "g_moyens-a_terre": "Accès à la Terre",
    "g_moyens-superficie_ha": "Superficie (ha)",
    "g_moyens-petits_ruminants": "Petits Ruminants",
    "g_moyens-bovins": "Bovins",
    "g_moyens-camelins": "Camelins",
    "g_moyens-equins": "Équins",
    "g_moyens-anes": "Ânes",
    "g_moyens-volaille": "Volaille",
    "g_moyens-moyens_production": "Moyens de Production",
    "g_moyens-stock_semences": "Stock Semences",
    "g_moyens-tlu": "TLU",
    "g_moyens-a_moyen_prod": "A Moyen Production",
    "h_alimentation-marche_5km": "Marché <5km",
    "h_alimentation-temps_marche": "Temps au Marché",
    "h_alimentation-repas_adultes": "Repas Adultes/jour",
    "h_alimentation-repas_enfants": "Repas Enfants/jour",
    "i_rcsi-note_rcsi": "Note rCSI",
    "i_rcsi-csi1": "rCSI 1",
    "i_rcsi-csi2": "rCSI 2",
    "i_rcsi-csi3": "rCSI 3",
    "i_rcsi-csi4": "rCSI 4",
    "i_rcsi-csi5": "rCSI 5",
    "i_rcsi-rcsi": "Score rCSI",
    "i_rcsi-note_rcsi_res": "Résultat rCSI",
    "j_hhs-hhs1": "HHS 1",
    "j_hhs-hhs2": "HHS 2",
    "j_hhs-hhs3": "HHS 3",
    "j_hhs-hhs": "Score HHS",
    "j_hhs-note_hhs": "Niveau HHS",
    "k_logement-type_logement": "Type Logement",
    "k_logement-endette": "Endetté",
    "k_logement-niveau_dette": "Niveau Dette",
    "l_lcsi-note_lcsi": "Note LCSI",
    "l_lcsi-lcsi_stress1": "LCSI Stress 1",
    "l_lcsi-lcsi_stress2": "LCSI Stress 2",
    "l_lcsi-lcsi_crise1": "LCSI Crise 1",
    "l_lcsi-lcsi_crise2": "LCSI Crise 2",
    "l_lcsi-lcsi_urg1": "LCSI Urgence 1",
    "l_lcsi-lcsi_urg2": "LCSI Urgence 2",
    "l_lcsi-lcsi_max": "LCSI Max",
    "m_score-eff_terre_p": "Eff Terre Pauvre",
    "m_score-eff_terre_m": "Eff Terre Moyen",
    "m_score-sc_a_cat": "Score A (Catégorie)",
    "m_score-sc_b_rcsi": "Score B rCSI",
    "m_score-sc_b_hhs": "Score B HHS",
    "m_score-sc_b_lcsi": "Score B LCSI",
    "m_score-sc_b": "Score B Total",
    "m_score-sc_c_terre": "Score C Terre",
    "m_score-sc_c_tlu": "Score C TLU",
    "m_score-sc_c_prod": "Score C Production",
    "m_score-sc_c": "Score C Total",
    "m_score-sc_d_dep": "Score D Dépendance",
    "m_score-sc_d": "Score D Total",
    "m_score-score_total": "Score Total",
    "m_score-classification": "Classification (code)",
    "m_score-classification_lbl": "Classification",
    "m_score-flag_contradiction": "Flag Contradiction",
    "m_score-flag_vad": "Flag VAD",
    "m_score-propose_auto": "Proposition Auto",
    "m_score-propose_lbl": "Proposition Label",
    "m_score-priorite": "Priorité",
    "m_score-affiche_score": "Affichage Score",
    "m_score-note_exclu_nanti": "Note Exclusion Nanti",
    "m_score-note_contradiction": "Note Contradiction",
    "m_score-note_vad": "Note VAD",
    "n_validation-rappel_decision": "Rappel Décision",
    "n_validation-agent_confirme": "Agent Confirme",
    "n_validation-decision_corrigee": "Décision Corrigée",
    "n_validation-justification": "Justification",
    "n_validation-decision_finale": "Décision Finale",
    "n_validation-beneficiaire_final": "Bénéficiaire Final",
    "n_validation-affiche_finale": "Affichage Final",
    "meta-instanceID": "Instance ID",
    "KEY": "Clé",
    "SubmitterID": "ID Soumetteur",
    "SubmitterName": "Nom Soumetteur",
    "AttachmentsPresent": "Pièces Jointes Présentes",
    "AttachmentsExpected": "Pièces Jointes Attendues",
    "Status": "Statut",
    "ReviewState": "État Révision",
    "DeviceID": "ID Appareil (2)",
    "Edits": "Modifications",
    "FormVersion": "Version Formulaire",
}

# Styles
HEADER_FILL = PatternFill("solid", fgColor="003366")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
SUBHEADER_FILL = PatternFill("solid", fgColor="1F6699")
SUBHEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(bold=True, color="003366", size=13)
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
SECTION_FONT = Font(bold=True, color="1F3864", size=11)
ALT_FILL = PatternFill("solid", fgColor="EBF0FA")
BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
THICK_BORDER = Border(
    left=Side(style="medium", color="003366"),
    right=Side(style="medium", color="003366"),
    top=Side(style="medium", color="003366"),
    bottom=Side(style="medium", color="003366"),
)

def apply_header_style(cell, fill=None, font=None):
    cell.fill = fill or HEADER_FILL
    cell.font = font or HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER

def apply_data_style(cell, alt=False):
    if alt:
        cell.fill = ALT_FILL
    cell.alignment = Alignment(vertical="center", wrap_text=False)
    cell.border = BORDER

wb = Workbook()

# ============================================================
# FEUILLE 1 — DONNÉES BRUTES
# ============================================================
ws_data = wb.active
ws_data.title = "Données"

# Titre
ws_data.merge_cells(f"A1:{get_column_letter(len(df.columns))}1")
title_cell = ws_data["A1"]
title_cell.value = "DONNÉES DE CIBLAGE DES MÉNAGES — ADK OUL"
title_cell.font = Font(bold=True, color="003366", size=14)
title_cell.alignment = Alignment(horizontal="center", vertical="center")
title_cell.fill = PatternFill("solid", fgColor="D9E1F2")
ws_data.row_dimensions[1].height = 28

# En-têtes traduits
headers = [col_labels.get(c, c) for c in df.columns]
for col_idx, header in enumerate(headers, start=1):
    cell = ws_data.cell(row=2, column=col_idx, value=header)
    apply_header_style(cell)
ws_data.row_dimensions[2].height = 40

# Données
for row_idx, row in enumerate(df.itertuples(index=False), start=3):
    alt = (row_idx % 2 == 0)
    for col_idx, value in enumerate(row, start=1):
        cell = ws_data.cell(row=row_idx, column=col_idx, value=value if pd.notna(str(value)) and value != "nan" else "")
        apply_data_style(cell, alt)

# Largeurs de colonnes
col_widths = {
    "Date de soumission": 22, "Nom Chef de Ménage": 22, "Village": 18,
    "Commune": 18, "Département": 18, "Région": 14, "Classification": 20,
    "Décision Finale": 20, "Nom Soumetteur": 20,
}
for col_idx, col_name in enumerate(headers, start=1):
    col_letter = get_column_letter(col_idx)
    if col_name in col_widths:
        ws_data.column_dimensions[col_letter].width = col_widths[col_name]
    else:
        ws_data.column_dimensions[col_letter].width = 14

ws_data.freeze_panes = "A3"

# ============================================================
# FEUILLE 2 — STATISTIQUES
# ============================================================
ws_stats = wb.create_sheet("Statistiques")

# Colonnes utiles
col_commune = "a_village-commune"
col_dept = "a_village-departement"
col_region = "a_village-region"
col_village = "a_village-village"
col_sexe = "b_chef-cm_sexe"
col_classif = "m_score-classification_lbl"
col_beneficiaire = "n_validation-beneficiaire_final"
col_statut_menage = "d_statut-statut_menage"
col_source = "g_moyens-source_principale"
col_taille = "c_composition-taille_menage"
col_score = "m_score-score_total"
col_submitter = "SubmitterName"
col_date = "SubmissionDate"

# Nettoyage
df[col_date] = pd.to_datetime(df[col_date], errors="coerce")
df[col_score] = pd.to_numeric(df[col_score], errors="coerce")
df[col_taille] = pd.to_numeric(df[col_taille], errors="coerce")

row = 1

def write_title(ws, row, text):
    ws.merge_cells(f"A{row}:L{row}")
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(bold=True, color="FFFFFF", size=12)
    c.fill = PatternFill("solid", fgColor="003366")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 26
    return row + 1

def write_section(ws, row, text):
    ws.merge_cells(f"A{row}:L{row}")
    c = ws.cell(row=row, column=1, value=text)
    c.font = SECTION_FONT
    c.fill = SECTION_FILL
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22
    return row + 1

def write_table_header(ws, row, headers_list):
    for col_i, h in enumerate(headers_list, 1):
        c = ws.cell(row=row, column=col_i, value=h)
        apply_header_style(c, fill=PatternFill("solid", fgColor="1F6699"))
    ws.row_dimensions[row].height = 32
    return row + 1

def write_row(ws, row, values, alt=False):
    for col_i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=col_i, value=v)
        if alt:
            c.fill = ALT_FILL
        c.alignment = Alignment(vertical="center")
        c.border = BORDER
    return row + 1

# TITRE GÉNÉRAL
row = write_title(ws_stats, row, "RAPPORT DE STATISTIQUES — CIBLAGE MÉNAGES ADK OUL")
row += 1

# 0. RÉSUMÉ GLOBAL
row = write_section(ws_stats, row, "0. RÉSUMÉ GLOBAL")
total = len(df)
nb_communes = df[col_commune].nunique()
nb_villages = df[col_village].nunique()
nb_dept = df[col_dept].nunique()
nb_beneficiaires = (df[col_beneficiaire] == "1").sum()

row = write_table_header(ws_stats, row, ["Indicateur", "Valeur"])
for i, (label, val) in enumerate([
    ("Total soumissions", total),
    ("Nombre de régions", df[col_region].nunique()),
    ("Nombre de départements", nb_dept),
    ("Nombre de communes", nb_communes),
    ("Nombre de villages", nb_villages),
    ("Nombre de bénéficiaires finaux (=1)", nb_beneficiaires),
    ("Score moyen", round(df[col_score].mean(), 1) if df[col_score].notna().any() else "N/A"),
    ("Taille moyenne des ménages", round(df[col_taille].mean(), 1) if df[col_taille].notna().any() else "N/A"),
]):
    row = write_row(ws_stats, row, [label, val], alt=(i % 2 == 0))
row += 1

# 1. SOUMISSIONS PAR COMMUNE
row = write_section(ws_stats, row, "1. SOUMISSIONS PAR COMMUNE")
row = write_table_header(ws_stats, row, [
    "Région", "Département", "Commune",
    "Total", "Masculin", "Féminin",
    "Très Vulnérable", "Vulnérable", "Moy. Vulnérable", "Non Vulnérable",
    "Bénéficiaires", "% Bénéficiaires"
])

commune_groups = df.groupby([col_region, col_dept, col_commune])
commune_rows = []
for (region, dept, commune), g in commune_groups:
    total_c = len(g)
    masc = (g[col_sexe] == "masculin").sum()
    fem = (g[col_sexe] == "feminin").sum()
    tres_vuln = (g[col_classif] == "TRES VULNERABLE").sum()
    vuln = (g[col_classif] == "VULNERABLE").sum()
    moy_vuln = (g[col_classif] == "MOYENNEMENT VULNERABLE").sum()
    non_vuln = (g[col_classif] == "NON VULNERABLE").sum()
    benef = (g[col_beneficiaire] == "1").sum()
    pct = f"{round(100 * benef / total_c, 1)}%" if total_c > 0 else "0%"
    commune_rows.append((region, dept, commune, total_c, masc, fem, tres_vuln, vuln, moy_vuln, non_vuln, benef, pct))

commune_rows.sort(key=lambda x: (-x[3], x[2]))
for i, r in enumerate(commune_rows):
    row = write_row(ws_stats, row, list(r), alt=(i % 2 == 0))

# Total commune
row = write_row(ws_stats, row, [
    "TOTAL", "", "",
    sum(r[3] for r in commune_rows),
    sum(r[4] for r in commune_rows),
    sum(r[5] for r in commune_rows),
    sum(r[6] for r in commune_rows),
    sum(r[7] for r in commune_rows),
    sum(r[8] for r in commune_rows),
    sum(r[9] for r in commune_rows),
    sum(r[10] for r in commune_rows),
    ""
])
ws_stats.cell(row=row-1, column=1).font = Font(bold=True)
row += 1

# 2. SOUMISSIONS PAR VILLAGE
row = write_section(ws_stats, row, "2. SOUMISSIONS PAR VILLAGE (TOP 20)")
row = write_table_header(ws_stats, row, [
    "Village", "Commune", "Département",
    "Total", "Masculin", "Féminin",
    "Très Vulnérable", "Vulnérable", "Moy. Vulnérable", "Non Vulnérable",
    "Bénéficiaires", "Score Moyen"
])

village_groups = df.groupby([col_village, col_commune, col_dept])
village_rows = []
for (village, commune, dept), g in village_groups:
    total_v = len(g)
    masc = (g[col_sexe] == "masculin").sum()
    fem = (g[col_sexe] == "feminin").sum()
    tres_vuln = (g[col_classif] == "TRES VULNERABLE").sum()
    vuln = (g[col_classif] == "VULNERABLE").sum()
    moy_vuln = (g[col_classif] == "MOYENNEMENT VULNERABLE").sum()
    non_vuln = (g[col_classif] == "NON VULNERABLE").sum()
    benef = (g[col_beneficiaire] == "1").sum()
    score_moy = round(g[col_score].mean(), 1) if g[col_score].notna().any() else "N/A"
    village_rows.append((village, commune, dept, total_v, masc, fem, tres_vuln, vuln, moy_vuln, non_vuln, benef, score_moy))

village_rows.sort(key=lambda x: -x[3])
for i, r in enumerate(village_rows[:20]):
    row = write_row(ws_stats, row, list(r), alt=(i % 2 == 0))
row += 1

# 3. RÉPARTITION PAR CLASSIFICATION
row = write_section(ws_stats, row, "3. RÉPARTITION PAR CLASSIFICATION DE VULNÉRABILITÉ")
row = write_table_header(ws_stats, row, ["Classification", "Nb Ménages", "% du Total", "Bénéficiaires", "% Bénéficiaires"])
classif_counts = df[col_classif].value_counts()
classif_order = ["TRES VULNERABLE", "VULNERABLE", "MOYENNEMENT VULNERABLE", "NON VULNERABLE"]
for i, c in enumerate(classif_order):
    n = classif_counts.get(c, 0)
    pct = f"{round(100 * n / total, 1)}%"
    benef_c = ((df[col_classif] == c) & (df[col_beneficiaire] == "1")).sum()
    pct_b = f"{round(100 * benef_c / n, 1)}%" if n > 0 else "0%"
    row = write_row(ws_stats, row, [c, n, pct, benef_c, pct_b], alt=(i % 2 == 0))
row += 1

# 4. RÉPARTITION PAR SEXE
row = write_section(ws_stats, row, "4. RÉPARTITION PAR SEXE DU CHEF DE MÉNAGE")
row = write_table_header(ws_stats, row, ["Sexe", "Nb Ménages", "% du Total", "Bénéficiaires", "Score Moyen"])
for i, (sexe, label) in enumerate([("masculin", "Masculin"), ("feminin", "Féminin")]):
    g = df[df[col_sexe] == sexe]
    n = len(g)
    pct = f"{round(100 * n / total, 1)}%"
    benef_s = (g[col_beneficiaire] == "1").sum()
    score_s = round(g[col_score].mean(), 1) if g[col_score].notna().any() else "N/A"
    row = write_row(ws_stats, row, [label, n, pct, benef_s, score_s], alt=(i % 2 == 0))
row += 1

# 5. SOURCE PRINCIPALE DE REVENUS
row = write_section(ws_stats, row, "5. SOURCE PRINCIPALE DE REVENUS")
row = write_table_header(ws_stats, row, ["Source", "Nb Ménages", "% du Total"])
source_counts = df[col_source].value_counts()
source_labels = {
    "agriculture": "Agriculture",
    "elevage": "Élevage",
    "commerce": "Commerce",
    "pas_emploi": "Sans emploi",
    "salarie": "Salarié",
    "artisanat": "Artisanat",
    "aide": "Aide/Transferts",
    "autre": "Autre",
}
for i, (src, n) in enumerate(source_counts.items()):
    label = source_labels.get(src, src)
    pct = f"{round(100 * n / total, 1)}%"
    row = write_row(ws_stats, row, [label, n, pct], alt=(i % 2 == 0))
row += 1

# 6. SOUMISSIONS PAR ENQUÊTEUR
row = write_section(ws_stats, row, "6. SOUMISSIONS PAR ENQUÊTEUR")
row = write_table_header(ws_stats, row, ["Nom Enquêteur", "Nb Soumissions", "% du Total", "Bénéficiaires", "Score Moyen"])
submitter_groups = df.groupby(col_submitter)
submitter_rows = []
for name, g in submitter_groups:
    n = len(g)
    benef_e = (g[col_beneficiaire] == "1").sum()
    score_e = round(g[col_score].mean(), 1) if g[col_score].notna().any() else "N/A"
    submitter_rows.append((name, n, f"{round(100*n/total,1)}%", benef_e, score_e))
submitter_rows.sort(key=lambda x: -x[1])
for i, r in enumerate(submitter_rows):
    row = write_row(ws_stats, row, list(r), alt=(i % 2 == 0))
row += 1

# 7. STATUT DES MÉNAGES (Communauté hôte / Déplacé)
row = write_section(ws_stats, row, "7. STATUT DES MÉNAGES")
row = write_table_header(ws_stats, row, ["Statut", "Nb Ménages", "% du Total"])
statut_counts = df[col_statut_menage].value_counts()
for i, (statut, n) in enumerate(statut_counts.items()):
    pct = f"{round(100 * n / total, 1)}%"
    row = write_row(ws_stats, row, [statut, n, pct], alt=(i % 2 == 0))
row += 1

# Largeurs colonnes stats
for col_i, width in enumerate([22, 18, 20, 10, 12, 12, 16, 14, 18, 16, 14, 14], 1):
    ws_stats.column_dimensions[get_column_letter(col_i)].width = width

ws_stats.freeze_panes = "A2"

# ============================================================
# FEUILLE 3 — TABLEAU CROISÉ COMMUNE × CLASSIFICATION
# ============================================================
ws_pivot = wb.create_sheet("Pivot Commune")

pivot = df.pivot_table(
    index=col_commune,
    columns=col_classif,
    values="KEY",
    aggfunc="count",
    fill_value=0
).reset_index()

row_p = 1
ws_pivot.merge_cells(f"A{row_p}:G{row_p}")
c = ws_pivot.cell(row=row_p, column=1, value="TABLEAU CROISÉ : COMMUNE × CLASSIFICATION")
c.font = Font(bold=True, color="FFFFFF", size=12)
c.fill = PatternFill("solid", fgColor="003366")
c.alignment = Alignment(horizontal="center", vertical="center")
ws_pivot.row_dimensions[row_p].height = 26
row_p += 2

headers_p = ["Commune"] + list(pivot.columns[1:]) + ["TOTAL"]
for col_i, h in enumerate(headers_p, 1):
    c = ws_pivot.cell(row=row_p, column=col_i, value=h)
    apply_header_style(c, fill=PatternFill("solid", fgColor="1F6699"))
ws_pivot.row_dimensions[row_p].height = 32
row_p += 1

for i, prow in enumerate(pivot.itertuples(index=False)):
    vals = list(prow)
    total_row = sum(vals[1:])
    for col_i, v in enumerate(vals + [total_row], 1):
        c = ws_pivot.cell(row=row_p, column=col_i, value=v)
        if i % 2 == 0:
            c.fill = ALT_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
    row_p += 1

# Ligne totaux
totals = ["TOTAL"] + [int(pivot[c].sum()) for c in pivot.columns[1:]] + [int(pivot.iloc[:, 1:].sum().sum())]
for col_i, v in enumerate(totals, 1):
    c = ws_pivot.cell(row=row_p, column=col_i, value=v)
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="D9E1F2")
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center")

for col_i in range(1, len(headers_p) + 1):
    ws_pivot.column_dimensions[get_column_letter(col_i)].width = 22

# Sauvegarde
out_path = "/tmp/claude-0/-home-user-Odk/5a50d722-e734-53cc-9ba5-8ea4c49a6d17/scratchpad/adkoul_menage_ciblage.xlsx"
wb.save(out_path)
print(f"Fichier généré : {out_path}")
print(f"Total lignes : {total} | Communes : {nb_communes} | Villages : {nb_villages}")
